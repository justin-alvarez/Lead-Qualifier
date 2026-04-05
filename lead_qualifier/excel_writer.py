"""Excel output with formatting using openpyxl."""

import json
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter


# Colors
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BLUE_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

COLUMNS = [
    ("Company Name", "name", 30),
    ("Domain", "domain", 25),
    ("City", "city", 15),
    ("State", "state", 8),
    ("Phone", "phone", 15),
    ("Revenue", "revenue", 15),
    ("Employees", "employees", 12),
    ("Google Reviews", "google_review_count", 14),
    ("Google Rating", "google_rating", 12),
    ("Detected Platforms", "detected_platforms", 25),
    ("Service Type", "service_type", 15),
    ("AI Reasoning", "ai_reasoning", 40),
    ("Status", "status", 14),
    ("Rejection Reason", "rejection_reason", 35),
]


def _format_revenue(value) -> str:
    if value is None:
        return ""
    try:
        v = float(value)
        if v >= 1_000_000_000:
            return f"${v/1_000_000_000:,.1f}B"
        if v >= 1_000_000:
            return f"${v/1_000_000:,.1f}M"
        return f"${v:,.0f}"
    except (ValueError, TypeError):
        return str(value)


def _write_sheet(ws, companies: List[Dict], sheet_title: str):
    """Write a sheet with company data and formatting."""
    # Headers
    for col_idx, (header, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    # Data rows
    for row_idx, company in enumerate(companies, 2):
        for col_idx, (_, field, _) in enumerate(COLUMNS, 1):
            value = company.get(field)

            # Format revenue
            if field == "revenue":
                value = _format_revenue(value)

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(field == "ai_reasoning"))

            # Alternating row colors
            if row_idx % 2 == 0:
                cell.fill = ALT_ROW_FILL

        # Status coloring
        status = company.get("status", "")
        status_col = next(
            (i for i, (_, f, _) in enumerate(COLUMNS, 1) if f == "status"), None
        )
        if status_col:
            status_cell = ws.cell(row=row_idx, column=status_col)
            if status == "QUALIFIED":
                status_cell.fill = GREEN_FILL
            elif status == "MANUAL_REVIEW":
                status_cell.fill = YELLOW_FILL
            elif status == "REJECTED":
                status_cell.fill = RED_FILL

        # Highlight detected platforms
        platform_col = next(
            (i for i, (_, f, _) in enumerate(COLUMNS, 1) if f == "detected_platforms"), None
        )
        if platform_col:
            platform_cell = ws.cell(row=row_idx, column=platform_col)
            platforms_val = company.get("detected_platforms", "")
            if platforms_val:
                platform_cell.fill = BLUE_FILL


def _write_summary_sheet(ws, stats: Dict[str, Any], config: Dict[str, Any],
                         platform_counts: Dict[str, int]):
    """Write the summary sheet."""
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 30

    bold = Font(bold=True, size=12)
    section_font = Font(bold=True, size=14, color="4472C4")

    row = 1

    # Run Settings
    ws.cell(row=row, column=1, value="Run Settings").font = section_font
    row += 1
    settings = [
        ("Min Revenue", f"${config.get('min_revenue', 0):,.0f}"),
        ("Min Employees", config.get("min_employees", 0)),
        ("Min Reviews", config.get("min_reviews", 0)),
        ("Service Type", config.get("service_type", "residential")),
        ("Database", config.get("db", "")),
    ]
    for label, value in settings:
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=str(value))
        row += 1

    row += 1

    # Funnel Stats
    ws.cell(row=row, column=1, value="Pipeline Funnel").font = section_font
    row += 1
    funnel = [
        ("Total in CSV", stats.get("total_in_csv", 0)),
        ("Skipped (duplicates)", stats.get("skipped_duplicates", 0)),
        ("New companies", stats.get("new_companies", 0)),
        ("Re-checked", stats.get("rechecked", 0)),
        ("Passed Gate 1", stats.get("passed_gate1", 0)),
        ("Passed Gate 2", stats.get("passed_gate2", 0)),
        ("Passed Gate 3", stats.get("passed_gate3", 0)),
        ("Qualified", stats.get("total_qualified", 0)),
        ("Rejected", stats.get("total_rejected", 0)),
        ("Manual Review", stats.get("total_manual_review", 0)),
    ]
    for label, value in funnel:
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1

    # Platform breakdown
    ws.cell(row=row, column=1, value="Platform Breakdown").font = section_font
    row += 1
    ws.cell(row=row, column=1, value="Platforms Detected").font = bold
    ws.cell(row=row, column=2, value=stats.get("platforms_detected", 0))
    row += 1
    for platform, count in sorted(platform_counts.items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=platform)
        ws.cell(row=row, column=2, value=count)
        row += 1

    row += 1

    # Cost
    ws.cell(row=row, column=1, value="Cost").font = section_font
    row += 1
    ws.cell(row=row, column=1, value="Estimated API Cost").font = bold
    ws.cell(row=row, column=2, value=f"${stats.get('estimated_cost', 0):.2f}")


def generate_excel(output_path: str, qualified: List[Dict], manual_review: List[Dict],
                   rejected: List[Dict], stats: Dict[str, Any], config: Dict[str, Any],
                   platform_counts: Optional[Dict[str, int]] = None):
    """Generate the formatted Excel workbook."""
    wb = Workbook()

    # Sheet 1: Qualified Leads
    ws_qualified = wb.active
    ws_qualified.title = "Qualified Leads"
    _write_sheet(ws_qualified, qualified, "Qualified Leads")

    # Sheet 2: Manual Review
    ws_manual = wb.create_sheet("Manual Review")
    _write_sheet(ws_manual, manual_review, "Manual Review")

    # Sheet 3: Rejected
    ws_rejected = wb.create_sheet("Rejected")
    _write_sheet(ws_rejected, rejected, "Rejected")

    # Sheet 4: Summary
    ws_summary = wb.create_sheet("Summary")
    _write_summary_sheet(ws_summary, stats, config, platform_counts or {})

    wb.save(output_path)


def generate_export_all_excel(output_path: str, companies: List[Dict]):
    """Generate Excel with all qualified leads from database."""
    wb = Workbook()
    ws = wb.active
    ws.title = "All Qualified Leads"
    _write_sheet(ws, companies, "All Qualified Leads")
    wb.save(output_path)
